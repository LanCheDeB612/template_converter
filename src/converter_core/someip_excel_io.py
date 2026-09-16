"""在 SOME/IP 标准模板副本中写入转换结果并保留模板结构。"""

from __future__ import annotations

from copy import deepcopy
from decimal import Decimal
import math
from numbers import Integral, Real
import os
from pathlib import Path
import re
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from converter_core.converters.someip.models import SomeipConversionData
from converter_core.excel_io import Formula

NODE_SHEET = "节点配置"
BEHAVIOR_SHEET = "通信行为"
INTERFACE_SHEET = "事件和方法"

NODE_HEADERS = ("节点名称", "IP地址", "Offer目标IP", "SD端口", "ETH端口")
BEHAVIOR_HEADERS = (
    "Service Name",
    "Service ID",
    "Instance ID",
    "Major Version",
    "Minor Version",
    "Offer TTL(s)",
    "Subscribe TTL(s)",
    "Subscribe ACK TTL(s)",
    "Offer Option数量",
    "Subscribe Option数量",
    "Subscribe ACK Option数量",
    "VLAN ID",
    "Transport Protocol",
    "Server",
    "Server IP",
    "Server Port UDP",
    "Server Port TCP",
    "EventgroupID",
    "Event Protocol",
    "Multicast IP",
    "Multicast Port",
    "Client",
    "Client IP",
    "Client Port UDP",
    "Client Port TCP",
)
INTERFACE_HEADERS = (
    "Service Name",
    "Service ID",
    "Method/Event Name",
    "Method/Event ID",
    "RPC Type",
    "UDP/TCP",
    "EventgroupID",
    "Cyclic Time (ms)",
)
SOMEIP_HEADERS = {
    NODE_SHEET: NODE_HEADERS,
    BEHAVIOR_SHEET: BEHAVIOR_HEADERS,
    INTERFACE_SHEET: INTERFACE_HEADERS,
}

_EXPECTED_SHEETS = tuple(SOMEIP_HEADERS)
_SHEET_DATA_PATTERN = re.compile(rb"<sheetData(?:\s[^>]*)?>.*?</sheetData>", re.DOTALL)
_DATA_VALIDATIONS_PATTERN = re.compile(
    rb"<dataValidations(?:\s[^>]*)?>.*?</dataValidations>", re.DOTALL
)
_CELL_XFS_PATTERN = re.compile(rb"<cellXfs(?:\s[^>]*)?>.*?</cellXfs>", re.DOTALL)
_DIMENSION_PATTERN = re.compile(rb'(<dimension\s+ref=")([A-Z]+)(\d+):([A-Z]+)(\d+)(")')
_CELL_COLUMN_PATTERN = re.compile(r"([A-Z]+)")
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_MAIN_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REVISION_NAMESPACE = "http://schemas.microsoft.com/office/spreadsheetml/2014/revision"
_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"
_OFFICE_RELATIONSHIP_ID = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
)
_PROTOCOL_VALIDATION = '"UDP,TCP,BOTH"'
_RPC_VALIDATION = (
    '"R/R Method,F/F Method,Event,Field Notify,Field Setter,Field Getter,'
    'Notification-fixed,Notification-changeable"'
)


def validate_someip_template(path: Path) -> None:
    """精确核对三个 Sheet 和 38 个表头，阻止写入错误版本的模板。"""
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if tuple(workbook.sheetnames) != _EXPECTED_SHEETS:
            raise ValueError(
                f"SOME/IP 标准模板 Sheet 必须为 {list(_EXPECTED_SHEETS)}，"
                f"实际为 {workbook.sheetnames}。"
            )
        for sheet_name, expected_headers in SOMEIP_HEADERS.items():
            sheet = workbook[sheet_name]
            actual_headers = tuple(
                sheet.cell(1, column).value
                for column in range(1, len(expected_headers) + 1)
            )
            if actual_headers != expected_headers:
                raise ValueError(
                    f"SOME/IP 标准模板“{sheet_name}”Sheet 的第 1 行表头与程序要求不一致。"
                )
            trailing_headers = tuple(
                sheet.cell(1, column).value
                for column in range(len(expected_headers) + 1, sheet.max_column + 1)
                if sheet.cell(1, column).value is not None
            )
            if trailing_headers:
                raise ValueError(
                    f"SOME/IP 标准模板“{sheet_name}”Sheet 存在未知表头："
                    f"{list(trailing_headers)}。"
                )
    finally:
        workbook.close()


def _worksheet_targets(archive: ZipFile) -> dict[str, str]:
    namespace = {"m": _MAIN_NAMESPACE}
    relation_namespace = {"r": _RELATIONSHIP_NAMESPACE}
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        relation.attrib["Id"]: relation.attrib["Target"]
        for relation in relationships.findall("r:Relationship", relation_namespace)
    }
    result = {}
    for sheet in workbook.findall("m:sheets/m:sheet", namespace):
        target = targets[sheet.attrib[_OFFICE_RELATIONSHIP_ID]].lstrip("/")
        result[sheet.attrib["name"]] = target if target.startswith("xl/") else f"xl/{target}"
    return result


def _row_two_styles(template: Path) -> dict[str, tuple[int, ...]]:
    """读取模板第 2 行的有效样式，供预格式区域以外的新行复用。"""
    workbook = load_workbook(template, read_only=False, data_only=False)
    try:
        return {
            sheet_name: tuple(
                workbook[sheet_name].cell(2, column).style_id
                for column in range(1, len(headers) + 1)
            )
            for sheet_name, headers in SOMEIP_HEADERS.items()
        }
    finally:
        workbook.close()


def _add_left_aligned_styles(
    xml: bytes,
    source_style_ids: set[int],
) -> tuple[bytes, dict[int, int]]:
    """复制模板数据样式并改为左对齐，数字仍按真实数值写入。"""
    match = _CELL_XFS_PATTERN.search(xml)
    if match is None:
        raise ValueError("SOME/IP 标准模板样式表缺少 cellXfs。")
    cell_xfs = ET.fromstring(match.group(0))
    original_styles = list(cell_xfs.findall("xf"))
    left_style_ids: dict[int, int] = {}
    for source_style_id in sorted(source_style_ids):
        if not 0 <= source_style_id < len(original_styles):
            raise ValueError(f"SOME/IP 标准模板包含无效样式 ID：{source_style_id}。")
        style = deepcopy(original_styles[source_style_id])
        alignment = style.find("alignment")
        if alignment is None:
            alignment = ET.SubElement(style, "alignment")
        alignment.set("horizontal", "left")
        alignment.set("vertical", "center")
        style.set("applyAlignment", "1")
        left_style_ids[source_style_id] = len(cell_xfs)
        cell_xfs.append(style)
    cell_xfs.set("count", str(len(cell_xfs)))
    replacement = ET.tostring(cell_xfs, encoding="utf-8", short_empty_elements=True)
    return xml[: match.start()] + replacement + xml[match.end() :], left_style_ids


def _cell_column(cell: ET.Element) -> int:
    coordinate = cell.attrib.get("r", "")
    match = _CELL_COLUMN_PATTERN.match(coordinate)
    if match is None:
        raise ValueError(f"无法识别单元格坐标 {coordinate!r}。")
    return column_index_from_string(match.group(1))


def _column_letter(column: int) -> str:
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _clear_cell_value(cell: ET.Element) -> None:
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)


def _set_cell_value(cell: ET.Element, value: object) -> None:
    _clear_cell_value(cell)
    if isinstance(value, Formula):
        formula = ET.SubElement(cell, "f")
        formula.text = value.expression
        ET.SubElement(cell, "v")
        return
    if isinstance(value, bool):
        cell.set("t", "b")
        ET.SubElement(cell, "v").text = "1" if value else "0"
        return
    if isinstance(value, Integral):
        ET.SubElement(cell, "v").text = str(value)
        return
    if isinstance(value, (Real, Decimal)):
        numeric_value = float(value)
        if not math.isfinite(numeric_value):
            raise ValueError(f"Excel 单元格不支持非有限数值 {value!r}。")
        ET.SubElement(cell, "v").text = str(value)
        return
    cell.set("t", "inlineStr")
    inline_string = ET.SubElement(cell, "is")
    text = ET.SubElement(inline_string, "t")
    rendered = str(value)
    if rendered != rendered.strip():
        text.set(_XML_SPACE, "preserve")
    text.text = rendered


def _patch_sheet_data(
    xml: bytes,
    headers: tuple[str, ...],
    rows_to_write: tuple[dict[str, object], ...],
    generated_cell_styles: tuple[int, ...],
) -> bytes:
    match = _SHEET_DATA_PATTERN.search(xml)
    if match is None:
        raise ValueError("SOME/IP 标准模板工作表缺少 sheetData。")
    sheet_data = ET.fromstring(match.group(0))
    rows = {int(row.attrib["r"]): row for row in sheet_data.findall("row")}

    # 输出只保留模板布局，不保留模板数据区中的示例值；随后仅写入本次转换生成的数据。
    for row_number, row in rows.items():
        if row_number < 2:
            continue
        for cell in row.findall("c"):
            if _cell_column(cell) <= len(headers):
                _clear_cell_value(cell)

    header_to_column = {header: column for column, header in enumerate(headers, start=1)}
    for row_number, values in enumerate(rows_to_write, start=2):
        row = rows.get(row_number)
        if row is None:
            row = ET.Element("row", {"r": str(row_number)})
            sheet_data.append(row)
            rows[row_number] = row
        cells = {_cell_column(cell): cell for cell in row.findall("c")}
        for header, value in values.items():
            if value is None or value == "":
                continue
            column = header_to_column[header]
            cell = cells.get(column)
            if cell is None:
                cell = ET.Element("c", {"r": f"{_column_letter(column)}{row_number}"})
                row.append(cell)
                cells[column] = cell
            cell.set("s", str(generated_cell_styles[column - 1]))
            _set_cell_value(cell, value)
        row[:] = sorted(row, key=_cell_column)

    sheet_data[:] = sorted(sheet_data, key=lambda item: int(item.attrib["r"]))
    replacement = ET.tostring(sheet_data, encoding="utf-8", short_empty_elements=True)
    patched = xml[: match.start()] + replacement + xml[match.end() :]
    maximum_written_row = len(rows_to_write) + 1

    def extend_dimension(dimension_match: re.Match[bytes]) -> bytes:
        current_row = int(dimension_match.group(5))
        if maximum_written_row <= current_row:
            return dimension_match.group(0)
        return b"".join(
            (
                dimension_match.group(1),
                dimension_match.group(2),
                dimension_match.group(3),
                b":",
                dimension_match.group(4),
                str(maximum_written_row).encode("ascii"),
                dimension_match.group(6),
            )
        )

    return _DIMENSION_PATTERN.sub(extend_dimension, patched, count=1)


def _fix_validations(xml: bytes, sheet_name: str) -> bytes:
    """把下拉应用到数据区，并在输出副本中补全 Notification 类型。"""
    match = _DATA_VALIDATIONS_PATTERN.search(xml)
    if match is None:
        raise ValueError(f"SOME/IP 标准模板“{sheet_name}”缺少数据验证。")
    ET.register_namespace("", _MAIN_NAMESPACE)
    ET.register_namespace("xr", _REVISION_NAMESPACE)
    wrapper = (
        f'<root xmlns="{_MAIN_NAMESPACE}" xmlns:xr="{_REVISION_NAMESPACE}">'.encode()
        + match.group(0)
        + b"</root>"
    )
    validations = ET.fromstring(wrapper)[0]
    validation_tag = f"{{{_MAIN_NAMESPACE}}}dataValidation"
    formula_tag = f"{{{_MAIN_NAMESPACE}}}formula1"
    items = list(validations.findall(validation_tag))

    if sheet_name == NODE_SHEET:
        expected = (("Ethernet1::Port1", "E2:E1048576", None),)
    elif sheet_name == BEHAVIOR_SHEET:
        expected = (("UDP,TCP,BOTH", "M2:M1048576 S2:S1048576", _PROTOCOL_VALIDATION),)
    else:
        expected = (
            ("UDP,TCP,BOTH", "F2:F1048576", _PROTOCOL_VALIDATION),
            ("R/R Method", "E2:E1048576", _RPC_VALIDATION),
        )

    for marker, target_range, replacement_formula in expected:
        matched = []
        for validation in items:
            formula = validation.find(formula_tag)
            if formula is not None and marker in (formula.text or ""):
                matched.append((validation, formula))
        if len(matched) != 1:
            raise ValueError(
                f"SOME/IP 标准模板“{sheet_name}”中的 {marker} 下拉数量不是 1。"
            )
        validation, formula = matched[0]
        if validation.attrib.get("type") != "list":
            raise ValueError(f"SOME/IP 标准模板“{sheet_name}”中的 {marker} 不是列表校验。")
        validation.set("sqref", target_range)
        if replacement_formula is not None:
            formula.text = replacement_formula

    validations.set("count", str(len(items)))
    replacement = ET.tostring(validations, encoding="utf-8", short_empty_elements=True)
    return xml[: match.start()] + replacement + xml[match.end() :]


def _node_ip_formula(node_cell: str, node_last_row: int) -> Formula:
    """按节点名称查找用户填写的 IP，未填写或未匹配时保持空白。"""
    lookup = (
        f"LOOKUP(2,1/('{NODE_SHEET}'!$A$2:$A${node_last_row}={node_cell}),"
        f"'{NODE_SHEET}'!$B$2:$B${node_last_row})"
    )
    # Excel 会把公式引用到的空 IP 显示为 0；先判断查找结果，避免用户误以为
    # 0 是转换器生成的有效地址。公式随节点配置修改自动更新，无需再次转换。
    return Formula(f'IFERROR(IF({lookup}=0,"",{lookup}),"")')


def _sheet_rows(data: SomeipConversionData) -> dict[str, tuple[dict[str, object], ...]]:
    node_rows = tuple({"节点名称": node_name} for node_name in data.node_names)
    node_last_row = len(data.node_names) + 1
    behavior_rows = tuple(
        {
            "Service Name": row.service_name,
            "Service ID": row.service_id,
            "Instance ID": row.instance_id,
            "Major Version": row.major_version,
            "Minor Version": row.minor_version,
            "Transport Protocol": row.transport_protocol,
            "Server": row.server,
            "Server IP": _node_ip_formula(f"N{output_row}", node_last_row),
            "Server Port UDP": row.server_port_udp,
            "Server Port TCP": row.server_port_tcp,
            "EventgroupID": row.eventgroup_id,
            "Event Protocol": row.event_protocol,
            "Client": row.client,
            "Client IP": _node_ip_formula(f"V{output_row}", node_last_row),
        }
        for output_row, row in enumerate(data.behavior_rows, start=2)
    )
    interface_rows = tuple(
        {
            "Service Name": row.service_name,
            "Service ID": row.service_id,
            "Method/Event Name": row.element_name,
            "Method/Event ID": row.element_id,
            "RPC Type": row.rpc_type,
            "UDP/TCP": row.protocol,
            "EventgroupID": row.eventgroup_id,
            "Cyclic Time (ms)": row.cycle_time,
        }
        for row in data.interface_rows
    )
    return {
        NODE_SHEET: node_rows,
        BEHAVIOR_SHEET: behavior_rows,
        INTERFACE_SHEET: interface_rows,
    }


def write_someip_template_copy(
    template: Path,
    output: Path,
    data: SomeipConversionData,
) -> None:
    """复制只读标准模板，仅替换三个业务 Sheet 的数据及输出样式。"""
    template = Path(template)
    output = Path(output)
    if template.resolve() == output.resolve():
        raise ValueError("输出文件不能覆盖只读 SOME/IP 标准模板。")
    validate_someip_template(template)
    rows_by_sheet = _sheet_rows(data)
    styles_by_sheet = _row_two_styles(template)
    output.parent.mkdir(parents=True, exist_ok=True)

    temporary_path: Path | None = None
    try:
        with NamedTemporaryFile(
            prefix=f".{output.name}.", suffix=".tmp", dir=output.parent, delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)
        with ZipFile(template, "r") as source:
            worksheet_paths = _worksheet_targets(source)
            source_style_ids = {
                style_id
                for sheet_name in _EXPECTED_SHEETS
                for style_id in styles_by_sheet[sheet_name]
            }
            patched_styles, left_style_ids = _add_left_aligned_styles(
                source.read("xl/styles.xml"), source_style_ids
            )
            replacements = {"xl/styles.xml": patched_styles}
            for sheet_name in _EXPECTED_SHEETS:
                worksheet_path = worksheet_paths[sheet_name]
                generated_styles = tuple(
                    left_style_ids[style_id]
                    for style_id in styles_by_sheet[sheet_name]
                )
                patched = _patch_sheet_data(
                    source.read(worksheet_path),
                    SOMEIP_HEADERS[sheet_name],
                    rows_by_sheet[sheet_name],
                    generated_styles,
                )
                patched = _fix_validations(patched, sheet_name)
                replacements[worksheet_path] = patched

            # ZIP 中未参与转换的条目逐字节复制，模板自身始终不被保存或改写。
            with ZipFile(
                temporary_path, "w", compression=ZIP_DEFLATED, allowZip64=True
            ) as destination:
                for entry in source.infolist():
                    destination.writestr(
                        entry,
                        replacements.get(entry.filename, source.read(entry.filename)),
                    )
        os.replace(temporary_path, output)
        temporary_path = None
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
