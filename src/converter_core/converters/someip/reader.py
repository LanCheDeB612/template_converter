"""按两行表头和业务块边界读取客户 SOME/IP 工作簿。"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
import math
from pathlib import Path
import re
from typing import NoReturn

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from converter_core.contracts import ConversionIssue
from converter_core.converters.someip.errors import SomeipConversionError
from converter_core.converters.someip.models import (
    EventgroupDefinition,
    InterfaceDefinition,
    ServiceDefinition,
    ServiceDeployment,
    SomeipSourceData,
)

DEPLOYMENT_SHEET = "Service_Deployment"
INTERFACE_SHEET = "Service_Interface"

_DEPLOYMENT_HEADERS = {
    "server": "Server ECU",
    "service_name": "Service Name",
    "service_id": "Service ID",
    "instance_id": "Service Instance ID",
    "major_version": "Major Version",
    "minor_version": "Minor Version",
    "server_port": "Server Port",
    "transport_protocol": "L4-Protocol",
    "clients": "Clients",
}
_INTERFACE_HEADERS = {
    "service_id": "Service ID",
    "service_name": "Service Name",
    "block_type": "MethodFieldEventgroup",
    "eventgroup_id": "Event group ID",
    "element_id": "Element ID",
    "element_name": "Name",
    "protocol": "L4-Pro",
    "cycle_time": "Cycle Time(ms)",
    "rpc_type": "TypeNotification-fixedNotification-changeableEventRR-InRR-OutFireForget",
    "clients": "Clients",
}
_SUPPORTED_EVENT_TYPES = {
    "event": "Event",
    "notification-fixed": "Notification-fixed",
    "notification-changeable": "Notification-changeable",
}


@dataclass
class _PendingMethod:
    """暂存 Method 定义，直到确认同一块内同时存在 RR-In 和 RR-Out。"""

    source_row: int
    service_id: int
    service_name: str
    element_name: str
    element_id: int
    protocol: str
    request_rows: list[int]
    response_rows: list[int]


def _text(value: object | None) -> str:
    return "" if value is None else str(value).strip()


def _header(value: object | None) -> str:
    """提取双语表头开头的英文部分，并忽略其中的排版空白。"""
    english_lines: list[str] = []
    for line in _text(value).splitlines():
        match = re.match(r"[\x00-\x7f]*", line)
        english = match.group(0).strip() if match is not None else ""
        if english:
            english_lines.append(english)
        if len(english) != len(line.strip()):
            break
    return "".join("".join(english_lines).split()).casefold()


def _fail(sheet: str, row: int | None, column: str, message: str) -> NoReturn:
    raise SomeipConversionError(
        ConversionIssue(
            level="error",
            message=message,
            sheet=sheet,
            row=row,
            column=column,
        )
    )


def _integer(value: object, *, sheet: str, row: int, column: str) -> int:
    try:
        if isinstance(value, bool):
            raise ValueError
        if isinstance(value, int):
            number = value
        elif isinstance(value, float) and value.is_integer():
            number = int(value)
        else:
            rendered = _text(value)
            number = int(rendered, 16) if rendered.casefold().startswith("0x") else int(rendered)
        if number < 0:
            raise ValueError
        return number
    except (TypeError, ValueError):
        _fail(
            sheet,
            row,
            column,
            f"{column} 必须是非负十进制整数或 0x 开头的十六进制整数，实际值为“{value}”。",
        )


def _cycle_time(value: object | None, *, row: int) -> int | float | None:
    if value is None or _text(value) == "":
        return None
    try:
        if isinstance(value, bool):
            raise ValueError
        number = Decimal(_text(value))
        rendered = float(number)
        if not math.isfinite(rendered) or number < 0:
            raise ValueError
        return int(number) if number == number.to_integral_value() else rendered
    except (InvalidOperation, ValueError):
        _fail(
            INTERFACE_SHEET,
            row,
            "Cycle Time (ms)",
            f"Cycle Time (ms) 必须是非负数，实际值为“{value}”。",
        )


def _column_map(sheet: Worksheet, expected: dict[str, str]) -> dict[str, int]:
    """按第 1 行的英文表头前缀定位列，避免依赖当前样表的固定列字母。"""
    actual = {cell.column: _header(cell.value) for cell in sheet[1] if _header(cell.value)}
    result: dict[str, int] = {}
    issues: list[ConversionIssue] = []
    for field, prefix in expected.items():
        target = _header(prefix)
        matches = [column for column, value in actual.items() if value == target]
        if not matches:
            matches = [column for column, value in actual.items() if value.startswith(target)]
        if not matches:
            issues.append(
                ConversionIssue(
                    level="error",
                    message=f"缺少必需表头“{prefix}”。",
                    sheet=sheet.title,
                    row=1,
                    column=prefix,
                )
            )
        elif len(matches) > 1:
            issues.append(
                ConversionIssue(
                    level="error",
                    message=f"表头“{prefix}”匹配到多列，无法确定读取位置。",
                    sheet=sheet.title,
                    row=1,
                    column=prefix,
                )
            )
        else:
            result[field] = matches[0]
    if issues:
        raise SomeipConversionError(*issues)
    return result


def _client_columns(sheet: Worksheet, start_column: int) -> tuple[tuple[int, str], ...]:
    """读取 Clients 分组第 2 行的连续节点名，空列表示该分组结束。"""
    clients: list[tuple[int, str]] = []
    for column in range(start_column, sheet.max_column + 1):
        name = _text(sheet.cell(2, column).value)
        if not name:
            if clients:
                break
            continue
        clients.append((column, name))
    if not clients:
        _fail(sheet.title, 2, "Clients", "Clients 分组第 2 行没有节点名称。")
    return tuple(clients)


def _marked_clients(
    sheet: Worksheet,
    row: int,
    client_columns: tuple[tuple[int, str], ...],
) -> tuple[str, ...]:
    result: list[str] = []
    for column, name in client_columns:
        marker = _text(sheet.cell(row, column).value)
        if not marker:
            continue
        if marker.casefold() != "x":
            _fail(
                sheet.title,
                row,
                name,
                f"客户端标记只能为空或 x，实际值为“{marker}”。",
            )
        result.append(name)
    return tuple(result)


def _required_text(value: object | None, *, sheet: str, row: int, column: str) -> str:
    rendered = _text(value)
    if not rendered or rendered == "-":
        _fail(sheet, row, column, f"{column} 不能为空。")
    return rendered


def _protocol(value: object | None, *, sheet: str, row: int, column: str) -> str:
    """第一版只转换真实样表已确认的 UDP，其他协议保留为明确错误。"""
    rendered = _required_text(value, sheet=sheet, row=row, column=column).upper()
    if rendered != "UDP":
        _fail(
            sheet,
            row,
            column,
            f"当前版本只支持已确认的 UDP，实际值为“{value}”。",
        )
    return rendered


def _read_deployments(sheet: Worksheet) -> tuple[ServiceDeployment, ...]:
    columns = _column_map(sheet, _DEPLOYMENT_HEADERS)
    clients = _client_columns(sheet, columns["clients"])
    records: list[ServiceDeployment] = []
    last_relevant_column = clients[-1][0]
    for row in range(3, sheet.max_row + 1):
        if not any(sheet.cell(row, column).value is not None for column in range(1, last_relevant_column + 1)):
            continue
        records.append(
            ServiceDeployment(
                source_row=row,
                server=_required_text(
                    sheet.cell(row, columns["server"]).value,
                    sheet=sheet.title,
                    row=row,
                    column="Server ECU",
                ),
                service_name=_required_text(
                    sheet.cell(row, columns["service_name"]).value,
                    sheet=sheet.title,
                    row=row,
                    column="Service Name",
                ),
                service_id=_integer(
                    sheet.cell(row, columns["service_id"]).value,
                    sheet=sheet.title,
                    row=row,
                    column="Service ID",
                ),
                instance_id=_integer(
                    sheet.cell(row, columns["instance_id"]).value,
                    sheet=sheet.title,
                    row=row,
                    column="Service Instance ID",
                ),
                major_version=_integer(
                    sheet.cell(row, columns["major_version"]).value,
                    sheet=sheet.title,
                    row=row,
                    column="Major Version",
                ),
                minor_version=_integer(
                    sheet.cell(row, columns["minor_version"]).value,
                    sheet=sheet.title,
                    row=row,
                    column="Minor Version",
                ),
                server_port=_integer(
                    sheet.cell(row, columns["server_port"]).value,
                    sheet=sheet.title,
                    row=row,
                    column="Server Port",
                ),
                transport_protocol=_protocol(
                    sheet.cell(row, columns["transport_protocol"]).value,
                    sheet=sheet.title,
                    row=row,
                    column="L4-Protocol",
                ),
                clients=_marked_clients(sheet, row, clients),
            )
        )
    if not records:
        _fail(sheet.title, None, "Service_Deployment", "没有可转换的服务部署记录。")
    return tuple(records)


def _read_interfaces(
    sheet: Worksheet,
) -> tuple[
    tuple[ServiceDefinition, ...],
    tuple[EventgroupDefinition, ...],
    tuple[InterfaceDefinition, ...],
]:
    columns = _column_map(sheet, _INTERFACE_HEADERS)
    clients = _client_columns(sheet, columns["clients"])
    services: list[ServiceDefinition] = []
    eventgroups: list[EventgroupDefinition] = []
    interfaces: list[InterfaceDefinition] = []
    current_service: ServiceDefinition | None = None
    current_eventgroup: EventgroupDefinition | None = None
    pending_method: _PendingMethod | None = None

    def finish_method() -> None:
        nonlocal pending_method
        if pending_method is None:
            return
        if len(pending_method.request_rows) != 1 or len(pending_method.response_rows) != 1:
            _fail(
                sheet.title,
                pending_method.source_row,
                "Type",
                "Method 必须在同一方法块内各有一条 RR-In 和 RR-Out。",
            )
        interfaces.append(
            InterfaceDefinition(
                source_row=pending_method.source_row,
                service_id=pending_method.service_id,
                service_name=pending_method.service_name,
                element_name=pending_method.element_name,
                element_id=pending_method.element_id,
                rpc_type="R/R Method",
                protocol=pending_method.protocol,
                eventgroup_id=None,
                cycle_time=None,
            )
        )
        pending_method = None

    for row in range(3, sheet.max_row + 1):
        service_id_value = sheet.cell(row, columns["service_id"]).value
        if _text(service_id_value):
            finish_method()
            current_service = ServiceDefinition(
                source_row=row,
                service_id=_integer(
                    service_id_value,
                    sheet=sheet.title,
                    row=row,
                    column="Service ID",
                ),
                service_name=_required_text(
                    sheet.cell(row, columns["service_name"]).value,
                    sheet=sheet.title,
                    row=row,
                    column="Service Name",
                ),
            )
            services.append(current_service)
            current_eventgroup = None

        block_type = _text(sheet.cell(row, columns["block_type"]).value)
        rpc_type = _text(sheet.cell(row, columns["rpc_type"]).value)
        if not block_type and not rpc_type:
            continue
        if current_service is None:
            _fail(sheet.title, row, "Service ID", "接口记录之前没有所属服务。")

        if block_type:
            finish_method()
            key = block_type.casefold()
            if key == "eventgroup":
                current_eventgroup = EventgroupDefinition(
                    source_row=row,
                    service_id=current_service.service_id,
                    service_name=current_service.service_name,
                    eventgroup_id=_integer(
                        sheet.cell(row, columns["eventgroup_id"]).value,
                        sheet=sheet.title,
                        row=row,
                        column="Event group ID",
                    ),
                    protocol=_protocol(
                        sheet.cell(row, columns["protocol"]).value,
                        sheet=sheet.title,
                        row=row,
                        column="L4-Pro",
                    ),
                    clients=_marked_clients(sheet, row, clients),
                )
                eventgroups.append(current_eventgroup)
            elif key == "method":
                current_eventgroup = None
                pending_method = _PendingMethod(
                    source_row=row,
                    service_id=current_service.service_id,
                    service_name=current_service.service_name,
                    element_name=_required_text(
                        sheet.cell(row, columns["element_name"]).value,
                        sheet=sheet.title,
                        row=row,
                        column="Name",
                    ),
                    element_id=_integer(
                        sheet.cell(row, columns["element_id"]).value,
                        sheet=sheet.title,
                        row=row,
                        column="Element ID",
                    ),
                    protocol=_protocol(
                        sheet.cell(row, columns["protocol"]).value,
                        sheet=sheet.title,
                        row=row,
                        column="L4-Pro",
                    ),
                    request_rows=[],
                    response_rows=[],
                )
            else:
                _fail(
                    sheet.title,
                    row,
                    "Method/Field/Eventgroup",
                    f"当前版本不支持接口块类型“{block_type}”。",
                )

        if not rpc_type:
            continue
        rpc_key = rpc_type.casefold()
        if rpc_key in {"rr-in", "rr-out"}:
            if pending_method is None:
                _fail(
                    sheet.title,
                    row,
                    "Type",
                    f"{rpc_type} 不在 Method 块内。",
                )
            target = (
                pending_method.request_rows
                if rpc_key == "rr-in"
                else pending_method.response_rows
            )
            target.append(row)
            continue

        converted_type = _SUPPORTED_EVENT_TYPES.get(rpc_key)
        if converted_type is None:
            _fail(
                sheet.title,
                row,
                "Type",
                f"当前版本不支持 RPC Type“{rpc_type}”。",
            )
        if current_eventgroup is None:
            _fail(
                sheet.title,
                row,
                "Event group ID",
                f"{rpc_type} 不在 Eventgroup 块内。",
            )
        own_protocol = sheet.cell(row, columns["protocol"]).value
        protocol = (
            _protocol(
                own_protocol,
                sheet=sheet.title,
                row=row,
                column="L4-Pro",
            )
            if _text(own_protocol)
            else current_eventgroup.protocol
        )
        interfaces.append(
            InterfaceDefinition(
                source_row=row,
                service_id=current_service.service_id,
                service_name=current_service.service_name,
                element_name=_required_text(
                    sheet.cell(row, columns["element_name"]).value,
                    sheet=sheet.title,
                    row=row,
                    column="Name",
                ),
                element_id=_integer(
                    sheet.cell(row, columns["element_id"]).value,
                    sheet=sheet.title,
                    row=row,
                    column="Element ID",
                ),
                rpc_type=converted_type,
                protocol=protocol,
                eventgroup_id=current_eventgroup.eventgroup_id,
                cycle_time=_cycle_time(
                    sheet.cell(row, columns["cycle_time"]).value,
                    row=row,
                ),
            )
        )

    finish_method()
    if not services:
        _fail(sheet.title, None, "Service_Interface", "没有可转换的服务定义。")
    return tuple(services), tuple(eventgroups), tuple(interfaces)


def read_customer_workbook(path: Path) -> SomeipSourceData:
    """只读取客户部署表和接口表，不修改或保存客户工作簿。"""
    workbook = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    try:
        missing = [
            name
            for name in (DEPLOYMENT_SHEET, INTERFACE_SHEET)
            if name not in workbook.sheetnames
        ]
        if missing:
            raise SomeipConversionError(
                *(
                    ConversionIssue(
                        level="error",
                        message=f"缺少必需 Sheet“{name}”。",
                        sheet=name,
                    )
                    for name in missing
                )
            )
        deployments = _read_deployments(workbook[DEPLOYMENT_SHEET])
        services, eventgroups, interfaces = _read_interfaces(workbook[INTERFACE_SHEET])
        return SomeipSourceData(
            deployments=deployments,
            services=services,
            eventgroups=eventgroups,
            interfaces=interfaces,
        )
    finally:
        workbook.close()
