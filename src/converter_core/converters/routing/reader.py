"""按完整表头读取客户路由工作簿。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from converter_core.contracts import ConversionIssue
from converter_core.converters.routing.errors import RoutingConversionError
from converter_core.converters.routing.models import CustomerEndpoint, CustomerRoutingRow

ROUTING_SHEET = "Routing Table"

_COMMON_HEADERS = (
    "Gateway Ecu",
    "Gateway Type(PDU/Signal)",
    "signal",
    "Size of Signal",
)
_SOURCE_HEADERS = {
    "network_name": "Source Network",
    "can_type": "Source CAN Type",
    "tcp_udp": "Source TCP/UDP",
    "vlan_id": "Source VLAN ID",
    "source_ip": "Source Src IP Addr",
    "target_ip": "Source Tgt IP Addr",
    "source_port": "Source Src Port",
    "target_port": "Source Tgt Port",
    "frame_tx_mode": "Source Frame TxMode",
    "frame_period": "Source Frame Period/mRTI(ms)",
    "frame_id": "Source Frame ID(Hex)",
    "frame_length": "Source Frame Length(Byte)",
    "lin_delay": "Source LIN Delay(ms)",
    "lin_schedule_delay": "Source LIN Schedule Delay(ms)",
    "pdu_name": "Source PDU",
    "pdu_id": "Source PDU ID(Hex)",
    "pdu_triggering": "Source PDU Triggering",
    "pdu_period": "Source PDU Period/mRTI(ms)",
    "pdu_length": "Source PDU Length(Byte)",
    "signal_start_bit": "Source Signal Start bit",
    "byte_order": "Source Byte Order",
}
_TARGET_HEADERS = {
    "network_name": "Target Network",
    "can_type": "Target CAN Type",
    "tcp_udp": "Target TCP/UDP",
    "vlan_id": "Target VLAN ID",
    "source_ip": "Target Src IP Addr",
    "target_ip": "Target Tgt IP Addr",
    "source_port": "Target Src Port",
    "target_port": "Target Tgt Port",
    "frame_tx_mode": "Target Frame TxMode",
    "frame_period": "Target Frame Period/mRTI(ms)",
    "frame_id": "Target Frame ID(Hex)",
    "frame_length": "Target Frame Length(Byte)",
    "lin_delay": "Target LIN Delay(ms)",
    "lin_schedule_delay": "Target LIN Schedule Delay(ms)",
    "pdu_name": "Target PDU",
    "pdu_id": "Target PDU ID(Hex)",
    "pdu_triggering": "Target PDU Triggering",
    "pdu_period": "Target PDU Period/mRTI(ms)",
    "pdu_length": "Target PDU Length(Byte)",
    "signal_start_bit": "Destination Signal Start bit",
    "byte_order": "Target Byte Order",
}


def _text(value: object | None) -> str:
    return "" if value is None else str(value).strip()


def _normalised_header(value: object | None) -> str:
    """客户表头可能含空格或换行，匹配时忽略排版差异。"""
    return "".join(_text(value).split()).casefold()


def _column_map(sheet: Worksheet) -> dict[str, int]:
    expected = (*_COMMON_HEADERS, *_SOURCE_HEADERS.values(), *_TARGET_HEADERS.values())
    actual: dict[str, list[int]] = {}
    for cell in sheet[1]:
        key = _normalised_header(cell.value)
        if key:
            actual.setdefault(key, []).append(cell.column)

    issues: list[ConversionIssue] = []
    result: dict[str, int] = {}
    for header in expected:
        columns = actual.get(_normalised_header(header), [])
        if not columns:
            issues.append(
                ConversionIssue(
                    level="error",
                    message=f"缺少必需表头“{header}”。",
                    sheet=sheet.title,
                    row=1,
                    column=header,
                )
            )
        elif len(columns) > 1:
            issues.append(
                ConversionIssue(
                    level="error",
                    message=f"表头“{header}”重复出现，无法确定应读取哪一列。",
                    sheet=sheet.title,
                    row=1,
                    column=header,
                )
            )
        else:
            result[header] = columns[0]
    if issues:
        raise RoutingConversionError(*issues)
    return result


def _endpoint(sheet: Worksheet, row: int, columns: dict[str, int], headers: dict[str, str]) -> CustomerEndpoint:
    values = {field: sheet.cell(row, columns[header]).value for field, header in headers.items()}
    return CustomerEndpoint(**values)


def read_routing_rows(sheet: Worksheet) -> tuple[CustomerRoutingRow, ...]:
    """读取每条客户路由；空白不继承上一行，原始值交给转换层判断。"""
    columns = _column_map(sheet)
    records: list[CustomerRoutingRow] = []
    for row in range(2, sheet.max_row + 1):
        # 客户表可能预先格式化大量空行，只有实际含值的行才属于转换输入。
        if not any(sheet.cell(row, column).value is not None for column in range(1, sheet.max_column + 1)):
            continue
        records.append(
            CustomerRoutingRow(
                source_row=row,
                gateway_ecu=_text(sheet.cell(row, columns["Gateway Ecu"]).value) or None,
                gateway_type=_text(sheet.cell(row, columns["Gateway Type(PDU/Signal)"]).value) or None,
                signal_name=sheet.cell(row, columns["signal"]).value,
                signal_length=sheet.cell(row, columns["Size of Signal"]).value,
                source=_endpoint(sheet, row, columns, _SOURCE_HEADERS),
                target=_endpoint(sheet, row, columns, _TARGET_HEADERS),
            )
        )
    return tuple(records)


def read_customer_workbook(path: Path) -> tuple[CustomerRoutingRow, ...]:
    """只读取客户工作簿的 ``Routing Table``，不改动输入文件。"""
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        if ROUTING_SHEET not in workbook.sheetnames:
            raise RoutingConversionError(
                ConversionIssue(
                    level="error",
                    message=f"缺少必需 Sheet“{ROUTING_SHEET}”。",
                    sheet=ROUTING_SHEET,
                )
            )
        return read_routing_rows(workbook[ROUTING_SHEET])
    finally:
        workbook.close()
