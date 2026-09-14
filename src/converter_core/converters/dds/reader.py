"""按 Sheet 和业务表头读取客户 DDS 工作簿。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from converter_core.contracts import ConversionIssue
from converter_core.converters.dds.errors import DdsConversionError
from converter_core.converters.dds.models import ServiceDeployment, TopicDefinition

DEPLOYMENT_SHEET = "ServiceDeployment"
TOPIC_SHEET = "ServicesAndTopicDefinition"


def _text(value: object | None) -> str:
    return "" if value is None else str(value).strip()


def _header(value: object | None) -> str:
    """忽略客户表头中的空白和换行，同时保留字段文字用于准确匹配。"""
    return "".join(_text(value).split()).casefold()


def _integer(value: object, *, sheet: str, row: int, column: str) -> int:
    try:
        if isinstance(value, bool):
            raise ValueError
        if isinstance(value, int):
            return value
        if isinstance(value, float) and value.is_integer():
            return int(value)
        text = _text(value)
        return int(text, 16) if text.casefold().startswith("0x") else int(text)
    except (TypeError, ValueError):
        raise DdsConversionError(
            ConversionIssue(
                level="error",
                message=f"{column} 必须是十进制整数或 0x 开头的十六进制整数，实际值为“{value}”。",
                sheet=sheet,
                row=row,
                column=column,
            )
        ) from None


def _find_column(sheet: Worksheet, header_row: int, expected: str) -> int:
    target = _header(expected)
    for cell in sheet[header_row]:
        if _header(cell.value) == target:
            return cell.column
    raise DdsConversionError(
        ConversionIssue(
            level="error",
            message=f"缺少必需表头“{expected}”。",
            sheet=sheet.title,
            row=header_row,
            column=expected,
        )
    )


def _merged_value(sheet: Worksheet, row: int, column: int) -> object | None:
    cell = sheet.cell(row, column)
    if cell.value is not None:
        return cell.value
    for merged_range in sheet.merged_cells.ranges:
        if cell.coordinate in merged_range:
            return sheet.cell(merged_range.min_row, merged_range.min_col).value
    return None


def _find_grouped_column(
    sheet: Worksheet,
    *,
    group_row: int,
    leaf_row: int,
    group: str,
    leaf: str,
) -> int:
    for column in range(1, sheet.max_column + 1):
        if _header(_merged_value(sheet, group_row, column)) == _header(group) and _header(
            sheet.cell(leaf_row, column).value
        ) == _header(leaf):
            return column
    raise DdsConversionError(
        ConversionIssue(
            level="error",
            message=f"缺少必需表头“{group}／{leaf}”。",
            sheet=sheet.title,
            row=leaf_row,
            column=f"{group}／{leaf}",
        )
    )


def _client_columns(sheet: Worksheet) -> tuple[tuple[int, str], ...]:
    clients_column = None
    for cell in sheet[1]:
        if _header(cell.value) == "clients":
            clients_column = cell.column
            break
    if clients_column is None:
        raise DdsConversionError(
            ConversionIssue(
                level="error",
                message="缺少必需表头“Clients”。",
                sheet=sheet.title,
                row=1,
                column="Clients",
            )
        )

    end_column = clients_column
    coordinate = sheet.cell(1, clients_column).coordinate
    for merged_range in sheet.merged_cells.ranges:
        if coordinate in merged_range:
            end_column = merged_range.max_col
            break

    columns = []
    for column in range(clients_column, end_column + 1):
        name = _text(sheet.cell(2, column).value)
        if not name:
            raise DdsConversionError(
                ConversionIssue(
                    level="error",
                    message="Clients 分组中存在空的客户端名称。",
                    sheet=sheet.title,
                    row=2,
                    column=sheet.cell(2, column).column_letter,
                )
            )
        columns.append((column, name))
    return tuple(columns)


def read_service_deployments(sheet: Worksheet) -> tuple[ServiceDeployment, ...]:
    """读取服务部署；客户工作表的格式化空行不会被当成业务数据。"""
    server_column = _find_column(sheet, 1, "Server ECU\n服务提供方")
    service_id_column = _find_column(sheet, 1, "Service ID\n服务ID")
    domain_column = _find_column(sheet, 1, "Domain ID\n域ID")
    client_columns = _client_columns(sheet)

    deployments = []
    for row in range(3, sheet.max_row + 1):
        source_values = [sheet.cell(row, column).value for column in range(1, sheet.max_column + 1)]
        if not any(value is not None for value in source_values):
            continue

        service_id_value = sheet.cell(row, service_id_column).value
        server = _text(sheet.cell(row, server_column).value)
        domain_value = sheet.cell(row, domain_column).value
        if service_id_value is None and not server and domain_value is None:
            continue
        if service_id_value is None or not server or domain_value is None:
            raise DdsConversionError(
                ConversionIssue(
                    level="error",
                    message="服务部署行必须同时填写 Server ECU、Service ID 和 Domain ID。",
                    sheet=sheet.title,
                    row=row,
                )
            )

        clients = tuple(
            name
            for column, name in client_columns
            if _text(sheet.cell(row, column).value).casefold() == "x"
        )
        deployments.append(
            ServiceDeployment(
                source_row=row,
                service_id=_integer(
                    service_id_value,
                    sheet=sheet.title,
                    row=row,
                    column=sheet.cell(row, service_id_column).column_letter,
                ),
                domain_id=_integer(
                    domain_value,
                    sheet=sheet.title,
                    row=row,
                    column=sheet.cell(row, domain_column).column_letter,
                ),
                server=server,
                clients=clients,
            )
        )
    if not deployments:
        raise DdsConversionError(
            ConversionIssue(level="error", message="没有找到服务部署数据。", sheet=sheet.title)
        )
    return tuple(deployments)


def _normalise_choice(
    value: object | None,
    aliases: dict[str, str],
    *,
    field: str,
    sheet: str,
    row: int,
    column: str,
) -> str | None:
    text = _text(value)
    if not text:
        return None
    key = text.replace("-", "_").casefold()
    if key not in aliases:
        raise DdsConversionError(
            ConversionIssue(
                level="error",
                message=f"{field} 的值“{value}”没有已确认的转换规则。",
                sheet=sheet,
                row=row,
                column=column,
            )
        )
    return aliases[key]


def read_topic_definitions(sheet: Worksheet) -> tuple[TopicDefinition, ...]:
    """按三层表头读取 Topic，Service ID 空白行继承最近的服务起始行。"""
    service_id_column = _find_column(sheet, 3, "Service ID")
    direction_column = _find_column(sheet, 3, "IN/OUT")
    topic_column = _find_column(sheet, 3, "TopicName")
    history_column = _find_grouped_column(
        sheet, group_row=2, leaf_row=3, group="HISTORY", leaf="Kind"
    )
    depth_column = _find_grouped_column(
        sheet, group_row=2, leaf_row=3, group="HISTORY", leaf="Depth"
    )
    reliability_column = _find_grouped_column(
        sheet, group_row=2, leaf_row=3, group="RELIABILITY", leaf="Kind"
    )
    max_samples_column = _find_column(sheet, 3, "Max_samples")
    max_instances_column = _find_column(sheet, 3, "Max_instances")
    max_per_instance_column = _find_column(sheet, 3, "Max_Samples_Per_Instance")

    current_service_id = None
    topics = []
    for row in range(4, sheet.max_row + 1):
        service_id_value = sheet.cell(row, service_id_column).value
        if service_id_value is not None and _text(service_id_value):
            current_service_id = _integer(
                service_id_value,
                sheet=sheet.title,
                row=row,
                column=sheet.cell(row, service_id_column).column_letter,
            )

        topic_name = _text(sheet.cell(row, topic_column).value)
        if not topic_name:
            continue
        if current_service_id is None:
            raise DdsConversionError(
                ConversionIssue(
                    level="error",
                    message="TopicName 所在行之前没有可关联的 Service ID。",
                    sheet=sheet.title,
                    row=row,
                    column=sheet.cell(row, topic_column).column_letter,
                )
            )

        direction = _text(sheet.cell(row, direction_column).value).upper()
        if direction not in {"IN", "OUT"}:
            raise DdsConversionError(
                ConversionIssue(
                    level="error",
                    message=f"包含 TopicName 的行必须将 IN/OUT 填写为 IN 或 OUT，实际值为“{direction}”。",
                    sheet=sheet.title,
                    row=row,
                    column=sheet.cell(row, direction_column).column_letter,
                )
            )

        limits = [
            sheet.cell(row, max_samples_column).value,
            sheet.cell(row, max_instances_column).value,
            sheet.cell(row, max_per_instance_column).value,
        ]
        resource_limits = None
        if any(_text(value) for value in limits):
            resource_limits = ",".join(_text(value) for value in limits)

        topics.append(
            TopicDefinition(
                source_row=row,
                service_id=current_service_id,
                topic_name=topic_name,
                direction=direction,
                reliability=_normalise_choice(
                    sheet.cell(row, reliability_column).value,
                    {"reliable": "RELIABLE", "best_effort": "BEST_EFFORT"},
                    field="RELIABILITY／Kind",
                    sheet=sheet.title,
                    row=row,
                    column=sheet.cell(row, reliability_column).column_letter,
                ),
                history=_normalise_choice(
                    sheet.cell(row, history_column).value,
                    {"keep_last": "KEEP_LAST", "keep_all": "KEEP_ALL"},
                    field="HISTORY／Kind",
                    sheet=sheet.title,
                    row=row,
                    column=sheet.cell(row, history_column).column_letter,
                ),
                history_depth=sheet.cell(row, depth_column).value,
                resource_limits=resource_limits,
            )
        )
    if not topics:
        raise DdsConversionError(
            ConversionIssue(level="error", message="没有找到 TopicName 数据。", sheet=sheet.title)
        )
    return tuple(topics)


def read_customer_workbook(path: Path) -> tuple[tuple[ServiceDeployment, ...], tuple[TopicDefinition, ...]]:
    """只读取本次已确认参与 DDS 转换的两个客户 Sheet。"""
    workbook = load_workbook(path, data_only=True, read_only=False)
    missing = [name for name in (DEPLOYMENT_SHEET, TOPIC_SHEET) if name not in workbook.sheetnames]
    if missing:
        raise DdsConversionError(
            *(ConversionIssue(level="error", message=f"缺少必需 Sheet“{name}”。", sheet=name) for name in missing)
        )
    return (
        read_service_deployments(workbook[DEPLOYMENT_SHEET]),
        read_topic_definitions(workbook[TOPIC_SHEET]),
    )
