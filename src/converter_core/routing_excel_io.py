"""在路由标准模板副本中写入转换结果，并保留模板原生对象。"""

from __future__ import annotations

from copy import deepcopy
from collections.abc import Mapping, Sequence
from decimal import Decimal
from numbers import Integral, Real
from pathlib import Path
import math
import os
import re
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

CAN_PDU_SHEET = "CAN-PDU路由"
SIGNAL_SHEET = "信号路由"
ETH_PDU_SHEET = "ETH-PDU路由"
NETWORK_MAPPING_SHEET = "网段映射"

CAN_PDU_HEADERS = (
    "pdu_name",
    "pdu_routing_type",
    "src_network_name",
    "src_network_type",
    "src_pdu_header_id",
    "src_pdu_length",
    "des_network_name",
    "des_network_type",
    "des_pdu_header_id",
    "des_cycle",
    "des_pdu_length",
)

SIGNAL_HEADERS = (
    "sig_name",
    "sig_routing_type",
    "sig_length",
    "src_network_name",
    "src_network_type",
    "src_pdu_header_id",
    "src_pdu_length",
    "src_startbit",
    "src_startbit_type",
    "src_byte_order",
    "des_network_name",
    "des_network_type",
    "des_pdu_header_id",
    "des_cycle",
    "des_pdu_length",
    "des_startbit",
    "des_startbit_type",
    "des_byte_order",
)

ETH_PDU_HEADERS = (
    "pdu_name",
    "pdu_routing_type",
    "src_network_name",
    "src_network_type",
    "src_pdu_header_id",
    "src_pdu_length",
    "src_eth_vlan_id",
    "src_eth_vlan_priority",
    "src_eth_client_port",
    "src_eth_client_mac",
    "src_eth_client_ip",
    "src_eth_server_port",
    "src_eth_server_mac",
    "src_eth_server_ip",
    "des_network_name",
    "des_network_type",
    "des_pdu_header_id",
    "des_cycle",
    "des_pdu_length",
    "des_eth_vlan_id",
    "des_eth_vlan_priority",
    "des_eth_client_port",
    "des_eth_client_mac",
    "des_eth_client_ip",
    "des_eth_server_port",
    "des_eth_server_mac",
    "des_eth_server_ip",
)

NETWORK_MAPPING_HEADERS = ("NetworkName", "TestChannel")

ROUTING_HEADERS = {
    CAN_PDU_SHEET: CAN_PDU_HEADERS,
    SIGNAL_SHEET: SIGNAL_HEADERS,
    ETH_PDU_SHEET: ETH_PDU_HEADERS,
    NETWORK_MAPPING_SHEET: NETWORK_MAPPING_HEADERS,
}

_WRITABLE_SHEETS = (CAN_PDU_SHEET, SIGNAL_SHEET, ETH_PDU_SHEET)
_EXPECTED_SHEETS = (*_WRITABLE_SHEETS, NETWORK_MAPPING_SHEET)
_SHEET_DATA_PATTERN = re.compile(rb"<sheetData(?:\s[^>]*)?>.*?</sheetData>", re.DOTALL)
_DATA_VALIDATIONS_PATTERN = re.compile(
    rb"<dataValidations(?:\s[^>]*)?>.*?</dataValidations>", re.DOTALL
)
_CELL_XFS_PATTERN = re.compile(rb"<cellXfs(?:\s[^>]*)?>.*?</cellXfs>", re.DOTALL)
_DIMENSION_PATTERN = re.compile(rb'(<dimension\s+ref=")([A-Z]+)(\d+):([A-Z]+)(\d+)(")')
_CELL_COLUMN_PATTERN = re.compile(r"([A-Z]+)")
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_MAIN_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REVISION_NAMESPACE = "http://schemas.microsoft.com/office/spreadsheetml/2014/revision"
_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"
_OFFICE_RELATIONSHIP_ID = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
)
_ETH_ROUTING_FORMULA = '"Event,Cycle,Always,Never"'
_ROUTING_VALIDATION_RANGE = "B2:B1048576"


def validate_routing_template(path: Path) -> None:
    """精确校验四个 Sheet 及 58 个业务表头。"""
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if tuple(workbook.sheetnames) != _EXPECTED_SHEETS:
            raise ValueError(
                f"路由标准模板 Sheet 必须为 {list(_EXPECTED_SHEETS)}，"
                f"实际为 {workbook.sheetnames}。"
            )

        for sheet_name, expected_headers in ROUTING_HEADERS.items():
            sheet = workbook[sheet_name]
            actual_headers = tuple(
                sheet.cell(1, column).value for column in range(1, len(expected_headers) + 1)
            )
            if actual_headers != expected_headers:
                raise ValueError(
                    f"路由标准模板“{sheet_name}”Sheet 的第 1 行表头"
                    "与程序要求不一致。"
                )

            # 模板可能因图形或格式范围产生额外列；只有第 1 行的
            # 非空文本才是业务表头，因此要单独防止未知表头混入。
            trailing_headers = tuple(
                sheet.cell(1, column).value
                for column in range(len(expected_headers) + 1, sheet.max_column + 1)
                if sheet.cell(1, column).value is not None
            )
            if trailing_headers:
                raise ValueError(
                    f"路由标准模板“{sheet_name}”Sheet 存在未知表头："
                    f"{list(trailing_headers)}。"
                )
    finally:
        workbook.close()


def _worksheet_targets(archive: ZipFile) -> dict[str, str]:
    namespace = {"m": _MAIN_NAMESPACE}
    relation_namespace = {"r": _RELATIONSHIP_NAMESPACE}
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        relation.attrib["Id"]: relation.attrib["Target"]
        for relation in relationships.findall("r:Relationship", relation_namespace)
    }
    result = {}
    for sheet in workbook.findall("m:sheets/m:sheet", namespace):
        target = targets[sheet.attrib[_OFFICE_RELATIONSHIP_ID]].lstrip("/")
        result[sheet.attrib["name"]] = target if target.startswith("xl/") else f"xl/{target}"
    return result


def _row_two_styles(template: Path) -> dict[str, tuple[int, ...]]:
    """读取每列第 2 行的有效样式。

    部分样式来自列维度而非单元格的 ``s`` 属性，由 openpyxl
    计算有效 style_id 可避免新增行丢失这类样式。
    """
    workbook = load_workbook(template, read_only=False, data_only=False)
    try:
        return {
            sheet_name: tuple(
                workbook[sheet_name].cell(2, column).style_id
                for column in range(1, len(headers) + 1)
            )
            for sheet_name, headers in ROUTING_HEADERS.items()
        }
    finally:
        workbook.close()


def _add_left_aligned_cell_styles(
    xml: bytes,
    source_style_ids: set[int],
) -> tuple[bytes, dict[int, int]]:
    """复制数据区现有样式并增加左对齐，避免把数字改成文本来影响下游解析。"""
    match = _CELL_XFS_PATTERN.search(xml)
    if match is None:
        raise ValueError("路由标准模板样式表缺少 cellXfs。")

    cell_xfs = ET.fromstring(match.group(0))
    original_styles = list(cell_xfs.findall("xf"))
    left_style_ids: dict[int, int] = {}
    for source_style_id in sorted(source_style_ids):
        if not 0 <= source_style_id < len(original_styles):
            raise ValueError(f"路由标准模板包含无效样式 ID：{source_style_id}。")

        style = deepcopy(original_styles[source_style_id])
        alignment = style.find("alignment")
        if alignment is None:
            alignment = ET.SubElement(style, "alignment")
        alignment.set("horizontal", "left")
        style.set("applyAlignment", "1")
        left_style_ids[source_style_id] = len(cell_xfs)
        cell_xfs.append(style)

    cell_xfs.set("count", str(len(cell_xfs)))
    replacement = ET.tostring(cell_xfs, encoding="utf-8", short_empty_elements=True)
    return xml[: match.start()] + replacement + xml[match.end() :], left_style_ids


def _cell_column(cell: ET.Element) -> int:
    coordinate = cell.attrib.get("r", "")
    match = _CELL_COLUMN_PATTERN.match(coordinate)
    if match is None:
        raise ValueError(f"无法识别单元格坐标 {coordinate!r}。")
    return column_index_from_string(match.group(1))


def _column_letter(column: int) -> str:
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _clear_cell_value(cell: ET.Element) -> None:
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)


def _set_cell_value(cell: ET.Element, value: object) -> None:
    _clear_cell_value(cell)
    if isinstance(value, bool):
        cell.set("t", "b")
        ET.SubElement(cell, "v").text = "1" if value else "0"
        return
    if isinstance(value, Integral):
        ET.SubElement(cell, "v").text = str(value)
        return
    if isinstance(value, (Real, Decimal)):
        numeric_value = float(value)
        if not math.isfinite(numeric_value):
            raise ValueError(f"Excel 单元格不支持非有限数值 {value!r}。")
        ET.SubElement(cell, "v").text = str(value)
        return

    # 普通字符串统一写成 inlineStr，避免改动模板共享字符串表。
    cell.set("t", "inlineStr")
    inline_string = ET.SubElement(cell, "is")
    text = ET.SubElement(inline_string, "t")
    rendered = str(value)
    if rendered != rendered.strip():
        text.set(_XML_SPACE, "preserve")
    text.text = rendered


def _normalise_sheet_rows(
    sheet_rows: Mapping[str, Sequence[Mapping[str, object]]],
) -> dict[str, tuple[Mapping[str, object], ...]]:
    unknown_sheets = set(sheet_rows) - set(_EXPECTED_SHEETS)
    if unknown_sheets:
        raise ValueError(f"输出数据包含未知 Sheet：{sorted(unknown_sheets)}。")

    mapping_rows = tuple(sheet_rows.get(NETWORK_MAPPING_SHEET, ()))
    if mapping_rows:
        raise ValueError("“网段映射”由人工填写，转换结果不允许写入数据行。")

    result = {}
    for sheet_name in _WRITABLE_SHEETS:
        headers = ROUTING_HEADERS[sheet_name]
        rows = tuple(sheet_rows.get(sheet_name, ()))
        for offset, row in enumerate(rows, start=2):
            if not isinstance(row, Mapping):
                raise TypeError(f"“{sheet_name}”第 {offset} 行必须是字段映射。")
            unknown_headers = set(row) - set(headers)
            if unknown_headers:
                raise ValueError(
                    f"“{sheet_name}”第 {offset} 行包含未知字段："
                    f"{sorted(unknown_headers)}。"
                )
        result[sheet_name] = rows
    result[NETWORK_MAPPING_SHEET] = ()
    return result


def _patch_sheet_data(
    xml: bytes,
    headers: tuple[str, ...],
    rows_to_write: Sequence[Mapping[str, object]],
    generated_cell_styles: tuple[int, ...],
) -> bytes:
    match = _SHEET_DATA_PATTERN.search(xml)
    if match is None:
        raise ValueError("路由标准模板工作表缺少 sheetData。")

    sheet_data = ET.fromstring(match.group(0))
    rows = {int(row.attrib["r"]): row for row in sheet_data.findall("row")}

    # 只清理业务列的值，保留行属性、样式单元格及非业务区域，
    # 避免破坏模板中的图形锚点或预设布局。
    for row_number, row in rows.items():
        if row_number < 2:
            continue
        for cell in row.findall("c"):
            if _cell_column(cell) <= len(headers):
                _clear_cell_value(cell)

    header_to_column = {header: column for column, header in enumerate(headers, start=1)}
    for row_number, values in enumerate(rows_to_write, start=2):
        row = rows.get(row_number)
        if row is None:
            row = ET.Element("row", {"r": str(row_number)})
            sheet_data.append(row)
            rows[row_number] = row
        cells = {_cell_column(cell): cell for cell in row.findall("c")}

        for header, value in values.items():
            column = header_to_column[header]
            # None 和空字符串表示真正空白，不写成 inlineStr 空值。
            if value is None or value == "":
                continue
            cell = cells.get(column)
            if cell is None:
                cell = ET.Element("c", {"r": f"{_column_letter(column)}{row_number}"})
                row.append(cell)
                cells[column] = cell
            # Excel 的“常规”对齐会让数值靠右。生成值使用保留原格式的
            # 左对齐派生样式，既统一视觉效果，也继续保存为真实数值。
            cell.set("s", str(generated_cell_styles[column - 1]))
            _set_cell_value(cell, value)
        row[:] = sorted(row, key=_cell_column)

    sheet_data[:] = sorted(sheet_data, key=lambda item: int(item.attrib["r"]))
    replacement = ET.tostring(sheet_data, encoding="utf-8", short_empty_elements=True)
    patched = xml[: match.start()] + replacement + xml[match.end() :]
    maximum_written_row = len(rows_to_write) + 1

    def extend_dimension(dimension_match: re.Match[bytes]) -> bytes:
        current_row = int(dimension_match.group(5))
        if maximum_written_row <= current_row:
            return dimension_match.group(0)
        return b"".join(
            (
                dimension_match.group(1),
                dimension_match.group(2),
                dimension_match.group(3),
                b":",
                dimension_match.group(4),
                str(maximum_written_row).encode("ascii"),
                dimension_match.group(6),
            )
        )

    return _DIMENSION_PATTERN.sub(extend_dimension, patched, count=1)


def _fix_eth_routing_validation(xml: bytes) -> bytes:
    """将 ETH 路由类型下拉修正到数据区，并移除误放在表头的校验。"""
    match = _DATA_VALIDATIONS_PATTERN.search(xml)
    if match is None:
        raise ValueError("路由标准模板“ETH-PDU路由”缺少数据验证。")

    ET.register_namespace("", _MAIN_NAMESPACE)
    ET.register_namespace("xr", _REVISION_NAMESPACE)
    # dataValidations 片段原本从 worksheet 根节点继承命名空间；
    # 单独解析时用临时根节点补齐，避免 xr:uid 成为未绑定前缀。
    wrapper = (
        f'<root xmlns="{_MAIN_NAMESPACE}" xmlns:xr="{_REVISION_NAMESPACE}">'.encode()
        + match.group(0)
        + b"</root>"
    )
    validations = ET.fromstring(wrapper)[0]
    validation_tag = f"{{{_MAIN_NAMESPACE}}}dataValidation"
    formula_tag = f"{{{_MAIN_NAMESPACE}}}formula1"
    target_validation: ET.Element | None = None

    for validation in list(validations.findall(validation_tag)):
        ranges = validation.attrib.get("sqref", "").split()
        if "B1" in ranges:
            ranges.remove("B1")
            if not ranges:
                validations.remove(validation)
                continue
            validation.set("sqref", " ".join(ranges))

        if _ROUTING_VALIDATION_RANGE not in ranges:
            continue
        if target_validation is not None:
            raise ValueError("ETH 路由类型数据区存在重复校验。")
        if validation.attrib.get("type") != "list":
            raise ValueError("ETH 路由类型数据区不是列表校验。")
        target_validation = validation

    if target_validation is None:
        raise ValueError(
            f"ETH 路由类型缺少 {_ROUTING_VALIDATION_RANGE} 列表校验。"
        )
    formula = target_validation.find(formula_tag)
    if formula is None:
        formula = ET.SubElement(target_validation, formula_tag)
    formula.text = _ETH_ROUTING_FORMULA
    validations.set("count", str(len(validations.findall(validation_tag))))

    replacement = ET.tostring(validations, encoding="utf-8", short_empty_elements=True)
    return xml[: match.start()] + replacement + xml[match.end() :]


def write_routing_template_copy(
    template: Path,
    output: Path,
    sheet_rows: Mapping[str, Sequence[Mapping[str, object]]],
) -> None:
    """复制路由标准模板并写入三类路由数据。

    ``sheet_rows`` 的外层键为 Sheet 名，内层键为相应 Sheet 的表头名。
    未提供的业务 Sheet 输出为空；“网段映射”始终只保留表头。
    """
    template = Path(template)
    output = Path(output)
    if template.resolve() == output.resolve():
        raise ValueError("输出文件不能覆盖只读路由标准模板。")

    validate_routing_template(template)
    normalised_rows = _normalise_sheet_rows(sheet_rows)
    styles_by_sheet = _row_two_styles(template)
    output.parent.mkdir(parents=True, exist_ok=True)

    temporary_path: Path | None = None
    try:
        with NamedTemporaryFile(
            prefix=f".{output.name}.", suffix=".tmp", dir=output.parent, delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)

        with ZipFile(template, "r") as source:
            worksheet_paths = _worksheet_targets(source)
            source_style_ids = {
                style_id
                for sheet_name in _WRITABLE_SHEETS
                for style_id in styles_by_sheet[sheet_name]
            }
            patched_styles, left_style_ids = _add_left_aligned_cell_styles(
                source.read("xl/styles.xml"), source_style_ids
            )
            replacements = {"xl/styles.xml": patched_styles}
            for sheet_name in _EXPECTED_SHEETS:
                worksheet_path = worksheet_paths[sheet_name]
                generated_cell_styles = tuple(
                    left_style_ids[style_id]
                    for style_id in styles_by_sheet[sheet_name]
                )
                patched = _patch_sheet_data(
                    source.read(worksheet_path),
                    ROUTING_HEADERS[sheet_name],
                    normalised_rows[sheet_name],
                    generated_cell_styles,
                )
                if sheet_name == ETH_PDU_SHEET:
                    patched = _fix_eth_routing_validation(patched)
                replacements[worksheet_path] = patched

            # 逐个复制 ZIP 条目，只替换工作表数据和新增左对齐样式；
            # 列宽、图形、关系文件和其他非数据对象均保持模板原状。
            with ZipFile(
                temporary_path, "w", compression=ZIP_DEFLATED, allowZip64=True
            ) as destination:
                for entry in source.infolist():
                    destination.writestr(
                        entry, replacements.get(entry.filename, source.read(entry.filename))
                    )
        os.replace(temporary_path, output)
        temporary_path = None
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
