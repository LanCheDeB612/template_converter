"""在复制的标准模板中写入数据，同时保留模板的非数据对象。"""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
import re
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from converter_core.converters.dds.models import DdsConversionData

NODE_SHEET = "节点配置"
MATRIX_SHEET = "通信矩阵"

NODE_HEADERS = (
    "Participant名称",
    "IP地址",
    "VLAN",
    "MAC地址",
    "DATA(p)组播MAC",
    "DATA(p)组播IP",
    "DATA(p)组播端口",
    "Domain ID",
    "Protocol Version",
    "Vendor ID",
    "Lease Duration(s)",
    "用户数据单播端口最小值",
    "用户数据单播端口最大值",
    "元数据单播端口",
)

MATRIX_HEADERS = (
    "Participant名称",
    "IP地址",
    "Domain ID",
    "Topic名称",
    "Topic类型",
    "DP角色",
    "Qos Reliability",
    "Qos Durability",
    "Qos DeadLine period\n(ns)",
    "Qos History",
    "Qos History depth",
    "Qos Resource Limits\n(max_samples,max_instances,max_samples_per_instance)",
    "Qos Latency Budget duration (ns)",
)

_SHEET_DATA_PATTERN = re.compile(rb"<sheetData(?:\s[^>]*)?>.*?</sheetData>", re.DOTALL)
_DIMENSION_PATTERN = re.compile(rb'(<dimension\s+ref=")([A-Z]+)(\d+):([A-Z]+)(\d+)(")')
_CELL_XFS_PATTERN = re.compile(
    rb'(<cellXfs\b[^>]*\bcount=")(\d+)("[^>]*>)(.*?)(</cellXfs>)',
    re.DOTALL,
)
_CELL_COLUMN_PATTERN = re.compile(r"([A-Z]+)")
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_BLACK_STYLE_INDEX = 14
_BLACK_WRAP_STYLE_INDEX = 18


@dataclass(frozen=True)
class Formula:
    """区分需要写入 Excel 公式的字符串和值字符串。"""

    expression: str


def _normalise_header(value: object | None) -> str:
    return "".join(("" if value is None else str(value)).split())


def validate_standard_template(path: Path) -> None:
    """写入前核对 Sheet 和表头，避免把数据写进错误版本的模板。"""
    workbook = load_workbook(path, read_only=False, data_only=False)
    if workbook.sheetnames != [NODE_SHEET, MATRIX_SHEET]:
        raise ValueError(
            f"标准模板 Sheet 必须为 {[NODE_SHEET, MATRIX_SHEET]}，实际为 {workbook.sheetnames}。"
        )
    for sheet_name, expected in ((NODE_SHEET, NODE_HEADERS), (MATRIX_SHEET, MATRIX_HEADERS)):
        sheet = workbook[sheet_name]
        actual = tuple(sheet.cell(1, column).value for column in range(1, len(expected) + 1))
        if tuple(map(_normalise_header, actual)) != tuple(map(_normalise_header, expected)):
            raise ValueError(f"标准模板“{sheet_name}”Sheet 的第 1 行表头与程序要求不一致。")


def _worksheet_targets(archive: ZipFile) -> dict[str, str]:
    namespace = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    relation_namespace = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
    relationship_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        relation.attrib["Id"]: relation.attrib["Target"]
        for relation in relationships.findall("r:Relationship", relation_namespace)
    }
    result = {}
    for sheet in workbook.findall("m:sheets/m:sheet", namespace):
        target = targets[sheet.attrib[relationship_id]].lstrip("/")
        result[sheet.attrib["name"]] = target if target.startswith("xl/") else f"xl/{target}"
    return result


def _cell_column(cell: ET.Element) -> int:
    match = _CELL_COLUMN_PATTERN.match(cell.attrib["r"])
    if match is None:
        raise ValueError(f"无法识别单元格坐标 {cell.attrib['r']}。")
    return column_index_from_string(match.group(1))


def _set_cell(cell: ET.Element, value: object) -> None:
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)

    if isinstance(value, Formula):
        formula = ET.SubElement(cell, "f")
        formula.text = value.expression
        ET.SubElement(cell, "v")
        return
    if isinstance(value, bool):
        cell.set("t", "b")
        ET.SubElement(cell, "v").text = "1" if value else "0"
        return
    if isinstance(value, (int, float)):
        ET.SubElement(cell, "v").text = str(value)
        return

    cell.set("t", "inlineStr")
    inline_string = ET.SubElement(cell, "is")
    text = ET.SubElement(inline_string, "t")
    rendered = str(value)
    if rendered != rendered.strip():
        text.set(_XML_SPACE, "preserve")
    text.text = rendered


def _patch_sheet_data(
    xml: bytes,
    values_by_row: dict[int, dict[int, object]],
    *,
    default_styles_by_column: dict[int, int] | None = None,
    force_styles_by_column: dict[int, int] | None = None,
) -> bytes:
    match = _SHEET_DATA_PATTERN.search(xml)
    if match is None:
        raise ValueError("标准模板工作表缺少 sheetData。")
    sheet_data = ET.fromstring(match.group(0))
    rows = {int(row.attrib["r"]): row for row in sheet_data.findall("row")}

    default_styles_by_column = default_styles_by_column or {}
    force_styles_by_column = force_styles_by_column or {}
    for row_number, values in values_by_row.items():
        row = rows.get(row_number)
        if row is None:
            row = ET.Element("row", {"r": str(row_number)})
            sheet_data.append(row)
            rows[row_number] = row
        cells = {_cell_column(cell): cell for cell in row.findall("c")}
        for column, value in values.items():
            cell = cells.get(column)
            if cell is None:
                cell = ET.Element("c", {"r": f"{_column_letter(column)}{row_number}"})
                row.append(cell)
                cells[column] = cell
                style_id = default_styles_by_column.get(column)
                if style_id is not None:
                    cell.set("s", str(style_id))
            style_id = force_styles_by_column.get(column)
            if style_id is not None:
                cell.set("s", str(style_id))
            _set_cell(cell, value)
        row[:] = sorted(row, key=_cell_column)

    sheet_data[:] = sorted(sheet_data, key=lambda row: int(row.attrib["r"]))
    replacement = ET.tostring(sheet_data, encoding="utf-8", short_empty_elements=True)
    patched = xml[: match.start()] + replacement + xml[match.end() :]
    maximum_row = max(values_by_row, default=1)

    def extend_dimension(dimension_match: re.Match[bytes]) -> bytes:
        current_row = int(dimension_match.group(5))
        if maximum_row <= current_row:
            return dimension_match.group(0)
        return b"".join(
            (
                dimension_match.group(1),
                dimension_match.group(2),
                dimension_match.group(3),
                b":",
                dimension_match.group(4),
                str(maximum_row).encode("ascii"),
                dimension_match.group(6),
            )
        )

    return _DIMENSION_PATTERN.sub(extend_dimension, patched, count=1)


def _column_letter(column: int) -> str:
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _append_black_left_styles(xml: bytes) -> tuple[bytes, int, int]:
    """在输出副本中追加黑色左对齐样式，并返回普通及自动换行样式编号。"""
    match = _CELL_XFS_PATTERN.search(xml)
    if match is None:
        raise ValueError("标准模板样式表缺少 cellXfs。")

    declared_count = int(match.group(2))
    root = ET.fromstring(b"<root>" + match.group(4) + b"</root>")
    styles = list(root)
    if declared_count != len(styles):
        raise ValueError("标准模板样式表的 cellXfs 数量不一致。")
    if len(styles) <= _BLACK_WRAP_STYLE_INDEX:
        raise ValueError("标准模板缺少生成数据所需的黑色基础样式。")

    appended = []
    for base_index, wrap_text in (
        (_BLACK_STYLE_INDEX, False),
        (_BLACK_WRAP_STYLE_INDEX, True),
    ):
        style = deepcopy(styles[base_index])
        alignment = style.find("alignment")
        if alignment is None:
            alignment = ET.SubElement(style, "alignment")
        alignment.set("horizontal", "left")
        alignment.set("vertical", "center")
        if wrap_text:
            alignment.set("wrapText", "1")
        else:
            alignment.attrib.pop("wrapText", None)
        style.set("applyFont", "1")
        style.set("applyAlignment", "1")
        appended.append(ET.tostring(style, encoding="utf-8", short_empty_elements=True))

    replacement = b"".join(
        (
            match.group(1),
            str(declared_count + len(appended)).encode("ascii"),
            match.group(3),
            match.group(4),
            *appended,
            match.group(5),
        )
    )
    patched = xml[: match.start()] + replacement + xml[match.end() :]
    return patched, declared_count, declared_count + 1


def _history_validation_is_correct(xml: bytes) -> bytes:
    """输出副本使用正确的 KEEP_LAST，不改变数据校验的其他属性。"""
    return xml.replace(b"KEPP_LAST", b"KEEP_LAST")


def _node_values(data: DdsConversionData) -> dict[int, dict[int, object]]:
    return {
        row: {1: node.participant_name, 8: node.domain_id}
        for row, node in enumerate(data.nodes, start=2)
    }


def _matrix_values(data: DdsConversionData) -> dict[int, dict[int, object]]:
    node_last_row = len(data.nodes) + 1
    values = {}
    for row, record in enumerate(data.matrix_rows, start=2):
        ip_lookup = (
            f'LOOKUP(2,1/(\'{NODE_SHEET}\'!$A$2:$A${node_last_row}=A{row})/'
            f'(\'{NODE_SHEET}\'!$H$2:$H${node_last_row}=C{row}),'
            f'\'{NODE_SHEET}\'!$B$2:$B${node_last_row})'
        )
        # Excel 会把公式引用到的空单元格显示为 0；先判断查找结果，
        # 未填写节点 IP 时返回空字符串，填写后仍随节点配置自动更新。
        ip_formula = f'IFERROR(IF({ip_lookup}=0,"",{ip_lookup}),"")'
        row_values: dict[int, object] = {
            1: record.participant_name,
            2: Formula(ip_formula),
            3: record.domain_id,
            4: record.topic_name,
            6: record.role,
        }
        if record.reliability is not None:
            row_values[7] = record.reliability
        if record.history is not None:
            row_values[10] = record.history
        if record.history_depth is not None:
            row_values[11] = record.history_depth
        if record.resource_limits is not None:
            row_values[12] = record.resource_limits
        values[row] = row_values
    return values


def write_dds_template_copy(template: Path, output: Path, data: DdsConversionData) -> None:
    """复制模板包，写入业务数据及统一样式，保留绘图等其他原生对象。"""
    validate_standard_template(template)
    output.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(template, "r") as source:
        worksheet_paths = _worksheet_targets(source)
        styles_xml, black_left_style, black_left_wrap_style = _append_black_left_styles(
            source.read("xl/styles.xml")
        )
        node_styles = {1: black_left_style, 8: black_left_style}
        matrix_styles = {
            1: black_left_style,
            2: black_left_style,
            3: black_left_style,
            4: black_left_wrap_style,
            6: black_left_style,
            7: black_left_style,
            10: black_left_style,
            11: black_left_style,
            12: black_left_wrap_style,
        }
        replacements = {
            "xl/styles.xml": styles_xml,
            worksheet_paths[NODE_SHEET]: _patch_sheet_data(
                source.read(worksheet_paths[NODE_SHEET]),
                _node_values(data),
                # 文本和数字都强制使用黑色左对齐样式，避免 Excel 按数据类型
                # 自动选择不同对齐方式。
                default_styles_by_column=node_styles,
                force_styles_by_column=node_styles,
            ),
            worksheet_paths[MATRIX_SHEET]: _history_validation_is_correct(
                _patch_sheet_data(
                    source.read(worksheet_paths[MATRIX_SHEET]),
                    _matrix_values(data),
                    # 覆盖模板预置区域中的红绿字体样式，并让超出预置范围的
                    # 新单元格保持相同的黑色左对齐显示。
                    default_styles_by_column=matrix_styles,
                    force_styles_by_column=matrix_styles,
                )
            ),
        }
        with ZipFile(output, "w", compression=ZIP_DEFLATED, allowZip64=True) as destination:
            for entry in source.infolist():
                destination.writestr(entry, replacements.get(entry.filename, source.read(entry.filename)))
