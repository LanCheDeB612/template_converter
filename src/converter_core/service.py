"""转换任务统一入口。"""

from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Callable
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from converter_core.contracts import (
    ConversionCount,
    ConversionIssue,
    ConversionRequest,
    ConversionResult,
)
from converter_core.converters.dds.converter import build_conversion_data
from converter_core.converters.dds.errors import DdsConversionError
from converter_core.converters.dds.reader import read_customer_workbook
from converter_core.converters.routing import (
    RoutingConversionData,
    RoutingConversionError,
    build_conversion_data as build_routing_conversion_data,
    read_customer_workbook as read_routing_workbook,
)
from converter_core.converters.someip import (
    SomeipConversionData,
    SomeipConversionError,
    build_conversion_data as build_someip_conversion_data,
    read_customer_workbook as read_someip_workbook,
)
from converter_core.excel_io import MATRIX_SHEET, NODE_SHEET, write_dds_template_copy
from converter_core.output import available_output_path
from converter_core.resources import (
    dds_template_path,
    routing_template_path,
    someip_template_path,
)
from converter_core.routing_excel_io import (
    CAN_PDU_HEADERS,
    CAN_PDU_SHEET,
    ETH_PDU_HEADERS,
    ETH_PDU_SHEET,
    NETWORK_MAPPING_SHEET,
    ROUTING_HEADERS,
    SIGNAL_HEADERS,
    SIGNAL_SHEET,
    write_routing_template_copy,
)
from converter_core.someip_excel_io import (
    BEHAVIOR_SHEET,
    INTERFACE_SHEET,
    NODE_SHEET as SOMEIP_NODE_SHEET,
    SOMEIP_HEADERS,
    write_someip_template_copy,
)

ProgressCallback = Callable[[int, str], None]


class ConversionService:
    """隔离界面和协议转换逻辑，并安全发布各类模板转换结果。"""

    def convert(
        self,
        request: ConversionRequest,
        progress: ProgressCallback | None = None,
    ) -> ConversionResult:
        """按模板类型执行转换，成功验证临时文件后才发布最终结果。"""
        problem = self._validate_request(request)
        if problem is not None:
            return ConversionResult(success=False, issues=(problem,))

        if request.template_type == "DDS":
            return self._convert_dds(request, progress)
        if request.template_type == "SOME/IP":
            return self._convert_someip(request, progress)
        return self._convert_routing(request, progress)

    def _convert_dds(
        self,
        request: ConversionRequest,
        progress: ProgressCallback | None,
    ) -> ConversionResult:
        """执行 DDS 读取、业务转换、模板写入和结果核对。"""
        notify = progress or (lambda _value, _message: None)

        temporary_path: Path | None = None
        try:
            notify(10, "正在读取客户 DDS 模板")
            deployments, topics = read_customer_workbook(request.input_file)
            notify(35, "正在关联服务、Topic 和 Participant")
            data = build_conversion_data(deployments, topics)

            template = dds_template_path()
            if not template.is_file():
                return ConversionResult(
                    success=False,
                    issues=(ConversionIssue(level="error", message=f"找不到标准模板：{template}"),),
                )

            output_file = available_output_path(
                request.input_file,
                request.output_directory,
                request.output_filename,
                template_type="DDS",
            )
            with NamedTemporaryFile(
                prefix=".dds-converting-",
                suffix=".xlsx",
                dir=request.output_directory,
                delete=False,
            ) as temporary:
                temporary_path = Path(temporary.name)

            notify(60, "正在复制标准模板并写入转换数据")
            write_dds_template_copy(template, temporary_path, data)
            notify(85, "正在重新打开并检查转换结果")
            self._verify_dds_output(temporary_path, len(data.nodes), len(data.matrix_rows))
            temporary_path.replace(output_file)
            temporary_path = None
            notify(100, "DDS 转换完成")
            return ConversionResult(
                success=True,
                output_file=output_file,
                counts=(
                    ConversionCount(label="节点", value=len(data.nodes)),
                    ConversionCount(label="通信记录", value=len(data.matrix_rows)),
                ),
                node_count=len(data.nodes),
                matrix_count=len(data.matrix_rows),
            )
        except DdsConversionError as error:
            return ConversionResult(success=False, issues=error.issues)
        except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError) as error:
            return ConversionResult(
                success=False,
                issues=(ConversionIssue(level="error", message=f"DDS 转换失败：{error}"),),
            )
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)

    def _convert_routing(
        self,
        request: ConversionRequest,
        progress: ProgressCallback | None,
    ) -> ConversionResult:
        """执行路由分流、字段转换、模板写入和结果核对。"""
        notify = progress or (lambda _value, _message: None)
        temporary_path: Path | None = None
        try:
            notify(10, "正在读取客户路由模板")
            rows = read_routing_workbook(request.input_file)
            notify(35, "正在识别协议方向并转换路由字段")
            data = build_routing_conversion_data(rows)

            template = routing_template_path()
            if not template.is_file():
                return ConversionResult(
                    success=False,
                    issues=(ConversionIssue(level="error", message=f"找不到标准模板：{template}"),),
                )

            output_file = available_output_path(
                request.input_file,
                request.output_directory,
                request.output_filename,
                template_type="路由",
            )
            with NamedTemporaryFile(
                prefix=".routing-converting-",
                suffix=".xlsx",
                dir=request.output_directory,
                delete=False,
            ) as temporary:
                temporary_path = Path(temporary.name)

            sheet_rows = self._routing_sheet_rows(data)
            notify(60, "正在复制标准模板并写入路由数据")
            write_routing_template_copy(template, temporary_path, sheet_rows)
            notify(85, "正在重新打开并检查路由转换结果")
            expected_counts = (
                len(data.can_pdu_routes),
                len(data.signal_routes),
                len(data.eth_pdu_routes),
            )
            self._verify_routing_output(temporary_path, expected_counts)
            temporary_path.replace(output_file)
            temporary_path = None
            notify(100, "路由转换完成")
            return ConversionResult(
                success=True,
                output_file=output_file,
                counts=(
                    ConversionCount(label="CAN-PDU", value=expected_counts[0]),
                    ConversionCount(label="信号路由", value=expected_counts[1]),
                    ConversionCount(label="ETH-PDU", value=expected_counts[2]),
                ),
            )
        except RoutingConversionError as error:
            return ConversionResult(success=False, issues=error.issues)
        except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError) as error:
            return ConversionResult(
                success=False,
                issues=(ConversionIssue(level="error", message=f"路由转换失败：{error}"),),
            )
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)

    def _convert_someip(
        self,
        request: ConversionRequest,
        progress: ProgressCallback | None,
    ) -> ConversionResult:
        """执行 SOME/IP 读取、块级转换、模板写入和结果核对。"""
        notify = progress or (lambda _value, _message: None)
        temporary_path: Path | None = None
        try:
            notify(10, "正在读取客户 SOME/IP 模板")
            source = read_someip_workbook(request.input_file)
            notify(35, "正在关联服务、事件组、客户端和接口定义")
            data = build_someip_conversion_data(source)

            template = someip_template_path()
            if not template.is_file():
                return ConversionResult(
                    success=False,
                    issues=(ConversionIssue(level="error", message=f"找不到标准模板：{template}"),),
                )
            output_file = available_output_path(
                request.input_file,
                request.output_directory,
                request.output_filename,
                template_type="SOME/IP",
            )
            with NamedTemporaryFile(
                prefix=".someip-converting-",
                suffix=".xlsx",
                dir=request.output_directory,
                delete=False,
            ) as temporary:
                temporary_path = Path(temporary.name)

            notify(60, "正在复制标准模板并写入 SOME/IP 数据")
            write_someip_template_copy(template, temporary_path, data)
            notify(85, "正在重新打开并检查 SOME/IP 转换结果")
            self._verify_someip_output(temporary_path, data)
            temporary_path.replace(output_file)
            temporary_path = None
            notify(100, "SOME/IP 转换完成")
            return ConversionResult(
                success=True,
                output_file=output_file,
                counts=(
                    ConversionCount(label="通信行为", value=len(data.behavior_rows)),
                    ConversionCount(label="事件和方法", value=len(data.interface_rows)),
                ),
            )
        except SomeipConversionError as error:
            return ConversionResult(success=False, issues=error.issues)
        except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError) as error:
            return ConversionResult(
                success=False,
                issues=(ConversionIssue(level="error", message=f"SOME/IP 转换失败：{error}"),),
            )
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)

    @staticmethod
    def _validate_request(request: ConversionRequest) -> ConversionIssue | None:
        if request.template_type not in {"DDS", "SOME/IP", "路由"}:
            return ConversionIssue(level="error", message=f"{request.template_type} 转换功能尚未实现。")
        if not request.input_file.is_file() or request.input_file.suffix.casefold() != ".xlsx":
            return ConversionIssue(level="error", message="请选择有效的 .xlsx 输入文件。")
        if not request.output_directory.is_dir():
            return ConversionIssue(level="error", message="请选择有效的输出目录。")
        return None

    @staticmethod
    def _verify_dds_output(path: Path, expected_nodes: int, expected_matrix_rows: int) -> None:
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            if workbook.sheetnames != [NODE_SHEET, MATRIX_SHEET]:
                raise ValueError("输出文件的 Sheet 结构发生变化。")
            node_sheet = workbook[NODE_SHEET]
            matrix_sheet = workbook[MATRIX_SHEET]
            node_count = sum(
                1
                for row in range(2, node_sheet.max_row + 1)
                if node_sheet.cell(row, 1).value is not None
            )
            matrix_count = sum(
                1
                for row in range(2, matrix_sheet.max_row + 1)
                if matrix_sheet.cell(row, 1).value is not None
            )
            if node_count != expected_nodes or matrix_count != expected_matrix_rows:
                raise ValueError(
                    f"输出重新打开后的记录数不一致：节点 {node_count}/{expected_nodes}，"
                    f"通信矩阵 {matrix_count}/{expected_matrix_rows}。"
                )
        finally:
            workbook.close()

    @staticmethod
    def _routing_sheet_rows(
        data: RoutingConversionData,
    ) -> dict[str, tuple[dict[str, object], ...]]:
        """去掉仅用于冲突定位的元数据，只把标准模板字段交给写入层。"""
        groups = (
            (CAN_PDU_SHEET, CAN_PDU_HEADERS, data.can_pdu_routes),
            (SIGNAL_SHEET, SIGNAL_HEADERS, data.signal_routes),
            (ETH_PDU_SHEET, ETH_PDU_HEADERS, data.eth_pdu_routes),
        )
        return {
            sheet_name: tuple(
                {header: getattr(record, header) for header in headers}
                for record in records
            )
            for sheet_name, headers, records in groups
        }

    @staticmethod
    def _verify_routing_output(path: Path, expected_counts: tuple[int, int, int]) -> None:
        """重新打开输出，核对四个 Sheet、三类数量、人工页空白和 ETH 枚举。"""
        workbook = load_workbook(path, read_only=False, data_only=False)
        expected_sheets = list(ROUTING_HEADERS)
        try:
            if workbook.sheetnames != expected_sheets:
                raise ValueError("路由输出文件的 Sheet 结构发生变化。")

            actual_counts = tuple(
                sum(
                    1
                    for row in range(2, workbook[sheet_name].max_row + 1)
                    if any(
                        workbook[sheet_name].cell(row, column).value is not None
                        for column in range(
                            1,
                            len(ROUTING_HEADERS[sheet_name]) + 1,
                        )
                    )
                )
                for sheet_name in expected_sheets[:3]
            )
            if actual_counts != expected_counts:
                raise ValueError(
                    "路由输出重新打开后的记录数不一致："
                    f"CAN-PDU {actual_counts[0]}/{expected_counts[0]}，"
                    f"信号路由 {actual_counts[1]}/{expected_counts[1]}，"
                    f"ETH-PDU {actual_counts[2]}/{expected_counts[2]}。"
                )

            mapping_sheet = workbook[NETWORK_MAPPING_SHEET]
            if any(
                mapping_sheet.cell(row, column).value is not None
                for row in range(2, mapping_sheet.max_row + 1)
                for column in range(1, len(("NetworkName", "TestChannel")) + 1)
            ):
                raise ValueError("网段映射应保留为空，不能复制模板示例数据。")

            eth_sheet = workbook[ETH_PDU_SHEET]
            route_type_validations = [
                validation
                for validation in eth_sheet.data_validations.dataValidation
                if "B2" in validation.sqref
            ]
            if (
                len(route_type_validations) != 1
                or route_type_validations[0].formula1 != '"Event,Cycle,Always,Never"'
            ):
                raise ValueError("ETH-PDU路由的 pdu_routing_type 下拉选项不完整。")
        finally:
            workbook.close()

    @staticmethod
    def _verify_someip_output(path: Path, data: SomeipConversionData) -> None:
        """重新打开输出，核对节点、关联公式和两类业务数量。"""
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            if workbook.sheetnames != list(SOMEIP_HEADERS):
                raise ValueError("SOME/IP 输出文件的 Sheet 结构发生变化。")
            node_sheet = workbook[SOMEIP_NODE_SHEET]
            actual_node_names = tuple(
                node_sheet.cell(row, 1).value
                for row in range(2, node_sheet.max_row + 1)
                if node_sheet.cell(row, 1).value is not None
            )
            if actual_node_names != data.node_names:
                raise ValueError("节点配置中的节点名称与客户 SOME/IP 模板不一致。")
            if any(
                node_sheet.cell(row, column).value is not None
                for row in range(2, len(data.node_names) + 2)
                for column in range(2, len(SOMEIP_HEADERS[SOMEIP_NODE_SHEET]) + 1)
            ):
                raise ValueError("节点配置除节点名称外应保持空白供用户手工填写。")

            behavior_sheet = workbook[BEHAVIOR_SHEET]
            if any(
                behavior_sheet.cell(row, column).data_type != "f"
                for row in range(2, len(data.behavior_rows) + 2)
                for column in (15, 23)
            ):
                raise ValueError("通信行为的 Server IP 或 Client IP 未按节点配置生成公式。")

            expected_counts = (len(data.behavior_rows), len(data.interface_rows))
            actual_counts = tuple(
                sum(
                    1
                    for row in range(2, workbook[sheet_name].max_row + 1)
                    if any(
                        workbook[sheet_name].cell(row, column).value is not None
                        for column in range(1, len(SOMEIP_HEADERS[sheet_name]) + 1)
                    )
                )
                for sheet_name in (BEHAVIOR_SHEET, INTERFACE_SHEET)
            )
            if actual_counts != expected_counts:
                raise ValueError(
                    "SOME/IP 输出重新打开后的记录数不一致："
                    f"通信行为 {actual_counts[0]}/{expected_counts[0]}，"
                    f"事件和方法 {actual_counts[1]}/{expected_counts[1]}。"
                )
        finally:
            workbook.close()
